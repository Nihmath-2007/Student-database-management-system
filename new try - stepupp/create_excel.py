import openpyxl
import os

STUDENT_DATA = [
    ("911524205001", "ABICHETHRA P", "07-08-2007"),
    ("911524205002", "ABIRAMI K", "21-04-2007"),
    ("911524205003", "AFRIN JAISHA S", "10-05-2007"),
    ("911524205004", "AHAMED RASIK FARITH S", "06-01-2006"),
    ("911524205005", "AHAMED SHAKI Z", "23-04-2005"),
    ("911524205006", "AL AFRAAH S", "30-10-2005"),
    ("911524205007", "ALBUNVINOJ D", "05-09-2006"),
    ("911524205008", "APRITH LENAN R J", "10-07-2006"),
    ("911524205009", "ASLAN SHA A", "10-05-2004"),
    ("911524205010", "BAHIYA SRI K", "22-05-2007"),
    ("911524205011", "BALAGANESH V", "14-10-2006"),
    ("911524205012", "DEVIKA P", "17-12-2006"),
    ("911524205013", "EDINA SHERIN A", "31-05-2007"),
    ("911524205015", "FAHIMA PARVEEN W", "18-09-2006"),
    ("911524205016", "FATHIMA FARHA M", "06-10-2006"),
    ("911524205017", "HARISH KUMAR R", "21-08-2006"),
    ("911524205018", "HARSHINI R", "13-03-2006"),
    ("911524205019", "ISMAIL D", "03-03-2006"),
    ("911524205020", "JEEVA N", "09-08-2006"),
    ("911524205021", "KALINIRANJAN K", "11-02-2007"),
    ("911524205022", "KARAN M", "24-12-2005"),
    ("911524205023", "KEERTHIGA M", "06-12-2006"),
    ("911524205024", "LAKSHMI NARAYANAN", "24-08-2005"),
    ("911524205025", "LITHISH K", "30-12-2006"),
    ("911524205026", "MATHANBABU M", "14-08-2007"),
    ("911524205027", "MOHAMED AATHIF H", "05-11-2005"),
    ("911524205028", "MOHAMED FAREETH K", "23-04-2007"),
    ("911524205029", "MOHAMED HAJA NAWAS J", "16-03-2007"),
    ("911524205030", "MOHAMED NIHMATHULLAH S", "03-03-2007"),
    ("911524205031", "MOHAMED SAJEER S", "20-04-2006"),
    ("911524205032", "MOHAMED SHEEDHU N", "15-04-2007"),
    ("911524205033", "MOHAMED THAMEEM ANSARI S", "18-06-2007"),
    ("911524205034", "MUNEESWARAN S", "09-11-2006"),
    ("911524205035", "NAAZIR MEERAN K", "10-09-2006"),
    ("911524205036", "NANDHINI R", "14-01-2007"),
    ("911524205037", "NIKILAN A", "21-07-2007"),
    ("911524205038", "NOORUL THAHANI K", "16-08-2007"),
    ("911524205039", "PAZHANI LAVANYA S", "09-07-2007"),
    ("911524205040", "PRITHIVIRAJ S", "09-01-2007"),
    ("911524205041", "RAHUL SANJAY R", "17-06-2007"),
    ("911524205042", "ROHITH RAJA A", "13-01-2007"),
    ("911524205043", "SAFRAN AATHIKA H", "15-12-2006"),
    ("911524205044", "SANJEEVPRAKASH S", "14-03-2007"),
    ("911524205045", "SARAFATH AHMED S", "14-03-2007"),
    ("911524205046", "SHABIYA THABASUM M", "21-08-2007"),
    ("911524205047", "SHAFAANA KHAALIDA S", "30-09-2006"),
    ("911524205048", "SHAKILA SRI T", "17-11-2006"),
    ("911524205050", "SIBI SHARVESH M", "10-09-2006"),
    ("911524205051", "SIFAUL ZABEEN M", "21-08-2006"),
    ("911524205052", "SUMAIYA S", "18-09-2005"),
    ("911524205053", "SWARNIKA B", "02-12-2006"),
    ("911524205301", "ABRISH SANJAI R", "23-06-2006"),
    ("911524205302", "DEVNATH C", "22-08-2006"),
    ("911524205303", "KANAGAVEL J", "02-11-2005"),
    ("911524205304", "MANOJ KUMAR S", "11-08-2006"),
    ("911524205305", "MOHAMED ASLAM S M", "29-08-2004"),
    ("911524205306", "MOHAMED HAAFIL H", "02-02-2006"),
    ("911524205307", "MOHAMED JASIR H", "25-09-2006"),
    ("911524205308", "RUTHISKAVI K", "02-05-2007"),
    ("911524205309", "SAMIH ANSAR A", "26-09-2004"),
    ("911524205310", "SYED ASHAKKEEN A", "17-10-2005"),
    ("911524205311", "VISHWA S", "22-06-2006")
]

def generate_excel(filepath):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Students"
    
    # Header row
    headers = ["Roll Number", "Student Name", "Date of Birth (DD-MM-YYYY)"]
    ws.append(headers)
    
    # Append data rows - ensure Roll Number is stored explicitly as text
    for roll, name, dob in STUDENT_DATA:
        row_num = ws.max_row + 1
        ws.cell(row=row_num, column=1, value=str(roll))
        ws.cell(row=row_num, column=1).data_type = 's'  # Explicit string type
        ws.cell(row=row_num, column=2, value=str(name))
        ws.cell(row=row_num, column=3, value=str(dob))
        
    wb.save(filepath)
    print(f"Generated {filepath} with {len(STUDENT_DATA)} students in sheet 'Students'.")

if __name__ == '__main__':
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    target_1 = os.path.join(base_dir, 'student_database.xlsx')
    target_2 = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'student_database.xlsx')
    generate_excel(target_1)
    generate_excel(target_2)
