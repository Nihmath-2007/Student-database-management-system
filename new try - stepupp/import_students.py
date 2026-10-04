import os
import sys
import datetime
import openpyxl
import bcrypt

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from config.database import get_db_connection

def hash_password(password: str) -> str:
    """Hash password using bcrypt."""
    salt = bcrypt.gensalt(rounds=10)
    return bcrypt.hashpw(password.encode('utf-8'), salt).decode('utf-8')

def format_dob_to_password(dob_val) -> str:
    """Convert DOB to DDMMYYYY format."""
    if isinstance(dob_val, (datetime.date, datetime.datetime)):
        return dob_val.strftime("%d%m%Y")
    dob_str = str(dob_val).strip()
    for sep in ['-', '/', '.', ' ']:
        dob_str = dob_str.replace(sep, '')
    return dob_str

def import_students_from_excel(excel_path=None):
    if not excel_path:
        possible_paths = [
            os.path.join(BASE_DIR, 'student_database.xlsx'),
            os.path.join(os.path.dirname(BASE_DIR), 'student_database.xlsx')
        ]
        for p in possible_paths:
            if os.path.isfile(p):
                excel_path = p
                break

    if not excel_path or not os.path.isfile(excel_path):
        raise FileNotFoundError(f"Excel file 'student_database.xlsx' not found. Checked: {possible_paths}")

    print(f"Loading Excel file: {excel_path}", flush=True)
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb["Students"]
    header_row = [cell.value for cell in ws[1]]
    
    roll_col, name_col, dob_col = 1, 2, 3
    for idx, header in enumerate(header_row, start=1):
        if not header:
            continue
        h_str = str(header).strip().lower()
        if "roll" in h_str or "reg" in h_str:
            roll_col = idx
        elif "name" in h_str:
            name_col = idx
        elif "birth" in h_str or "dob" in h_str:
            dob_col = idx

    print("Parsing student records from sheet 'Students'...", flush=True)
    records = []
    for row in range(2, ws.max_row + 1):
        roll_val = ws.cell(row=row, column=roll_col).value
        name_val = ws.cell(row=row, column=name_col).value
        dob_val = ws.cell(row=row, column=dob_col).value
        
        if not roll_val or not name_val or not dob_val:
            continue
            
        roll_number = str(roll_val).strip()
        student_name = str(name_val).strip()
        dob_raw = str(dob_val).strip() if not isinstance(dob_val, (datetime.date, datetime.datetime)) else dob_val.strftime("%d-%m-%Y")
        password_plain = format_dob_to_password(dob_val)
        records.append((roll_number, student_name, dob_raw, password_plain))

    print(f"Parsed {len(records)} students. Generating bcrypt password hashes...", flush=True)
    hashed_records = []
    for roll, name, dob_raw, pw_plain in records:
        pw_hash = hash_password(pw_plain)
        hashed_records.append((roll, name, dob_raw, pw_hash))

    print("Connecting to database...", flush=True)
    conn, engine = get_db_connection()
    try:
        cursor = conn.cursor(dictionary=True) if engine == 'mysql' else conn.cursor()
        
        # Fetch current state
        cursor.execute("SELECT studentid, regno FROM students")
        existing_students = cursor.fetchall()
        if engine == 'sqlite':
            existing_students = [{'studentid': r[0], 'regno': r[1]} for r in existing_students]
        student_id_by_roll = {str(s['regno']).strip(): s['studentid'] for s in existing_students if s.get('regno')}

        cursor.execute("SELECT MAX(studentid) as max_id FROM students")
        row = cursor.fetchone()
        current_max_id = (row['max_id'] if engine == 'mysql' else row[0]) or 0

        cursor.execute("SELECT id, username, student_id FROM users")
        existing_users = cursor.fetchall()
        if engine == 'sqlite':
            existing_users = [{'id': r[0], 'username': r[1], 'student_id': r[2]} for r in existing_users]
        user_by_uname = {str(u['username']).strip(): u['id'] for u in existing_users if u.get('username')}
        user_by_stud_id = {u['student_id']: u['id'] for u in existing_users if u.get('student_id')}

        p = "%s" if engine == 'mysql' else "?"

        students_to_update = []
        students_to_insert = []
        users_to_update = []
        users_to_insert = []

        for roll, name, dob_raw, pw_hash in hashed_records:
            if roll in student_id_by_roll:
                s_id = student_id_by_roll[roll]
                students_to_update.append((name, dob_raw, s_id))
            else:
                current_max_id += 1
                s_id = current_max_id
                student_id_by_roll[roll] = s_id
                students_to_insert.append((s_id, roll, name, dob_raw))

            u_id = user_by_uname.get(roll) or user_by_stud_id.get(s_id)
            if u_id:
                users_to_update.append((roll, pw_hash, s_id, u_id))
            else:
                users_to_insert.append((roll, pw_hash, s_id))

        print(f"Executing batch operations: updating {len(students_to_update)} students, inserting {len(students_to_insert)} students...", flush=True)
        if students_to_update:
            cursor.executemany(f"UPDATE students SET name = {p}, dob = {p} WHERE studentid = {p}", students_to_update)
        if students_to_insert:
            cursor.executemany(f"INSERT INTO students (studentid, regno, name, dob, department, year, email, section) VALUES ({p}, {p}, {p}, {p}, 'Information Technology', 'Third year', '', 'A')", students_to_insert)

        print(f"Executing batch operations: updating {len(users_to_update)} users, inserting {len(users_to_insert)} users...", flush=True)
        if users_to_update:
            cursor.executemany(f"UPDATE users SET username = {p}, password_hash = {p}, role = 'student', student_id = {p}, failed_attempts = 0, locked_until = NULL WHERE id = {p}", users_to_update)
        if users_to_insert:
            cursor.executemany(f"INSERT INTO users (username, password_hash, role, student_id, staff_id, failed_attempts, locked_until) VALUES ({p}, {p}, 'student', {p}, NULL, 0, NULL)", users_to_insert)

        # Admin account
        admin_hash = hash_password('admin123')
        cursor.execute("SELECT id FROM users WHERE username = 'admin'")
        admin_row = cursor.fetchone()
        if admin_row:
            admin_id = admin_row['id'] if engine == 'mysql' else admin_row[0]
            cursor.execute(f"UPDATE users SET password_hash = {p}, role = 'admin', failed_attempts = 0, locked_until = NULL WHERE id = {p}", (admin_hash, admin_id))
        else:
            cursor.execute(f"INSERT INTO users (username, password_hash, role, student_id, staff_id, failed_attempts, locked_until) VALUES ('admin', {p}, 'admin', NULL, NULL, 0, NULL)", (admin_hash,))

        # HOD account unlocked
        cursor.execute(f"UPDATE users SET failed_attempts = 0, locked_until = NULL WHERE username = 'hod'")

        conn.commit()
        cursor.close()

        print("\n" + "="*50, flush=True)
        print("DATABASE IMPORT SUCCESSFULLY COMPLETED!", flush=True)
        print(f"Total students processed: {len(hashed_records)}", flush=True)
        print(f"Students inserted: {len(students_to_insert)}, updated: {len(students_to_update)}", flush=True)
        print(f"Users inserted: {len(users_to_insert)}, updated: {len(users_to_update)}", flush=True)
        print("Admin user: 'admin' | Password: 'admin123'", flush=True)
        print("="*50 + "\n", flush=True)

    finally:
        conn.close()

if __name__ == '__main__':
    import_students_from_excel()
