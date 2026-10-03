import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from config.database import fetch_all

def get_marks_data_filepaths():
    """Returns primary and mirrored file paths for marks_data.xlsx."""
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    parent_dir = os.path.dirname(base_dir)

    target_primary = os.path.join(base_dir, 'marks_data.xlsx')
    target_mirror = os.path.join(parent_dir, 'marks_data.xlsx') if os.path.isdir(parent_dir) else None

    return target_primary, target_mirror


def _apply_table_styling(ws):
    """Applies clean, institutional styling to the Marks worksheet."""
    header_fill = PatternFill(start_color="0B1A30", end_color="0B1A30", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Arial", size=10, color="22252A")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # Style Header Row
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[1].height = 26

    # Style Data Rows
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        ws.row_dimensions[row[0].row].height = 20
        for idx, cell in enumerate(row):
            cell.font = data_font
            cell.border = thin_border
            # Alignments
            if idx in (0, 3):  # Reg No, Test Number
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif idx in (4, 5):  # Marks Obtained, Max Marks
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Auto-adjust column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 14)


def update_marks_data_file(records):
    """
    Updates or regenerates marks_data.xlsx (sheet 'Marks') with columns:
    [Register No, Student Name, Subject, Test Number, Marks Obtained, Max Marks]
    Updates existing rows if (Register No, Subject, Test Number) matches; does not duplicate them.
    """
    if not records:
        return

    primary_path, mirror_path = get_marks_data_filepaths()

    # Load existing workbook or create new
    if os.path.isfile(primary_path):
        try:
            wb = openpyxl.load_workbook(primary_path)
        except Exception:
            wb = openpyxl.Workbook()
    else:
        wb = openpyxl.Workbook()

    # Obtain or create sheet 'Marks'
    if "Marks" in wb.sheetnames:
        ws = wb["Marks"]
    else:
        if "Sheet" in wb.sheetnames and len(wb.sheetnames) == 1:
            ws = wb["Sheet"]
            ws.title = "Marks"
        else:
            ws = wb.create_sheet("Marks")

    headers = ["Register No", "Student Name", "Subject", "Test Number", "Marks Obtained", "Max Marks"]

    # Ensure header row exists
    if ws.max_row == 0 or not ws.cell(row=1, column=1).value:
        ws.delete_rows(1, ws.max_row)
        ws.append(headers)

    # Index existing rows by composite key: (Register No, Subject, Test Number)
    existing_key_map = {}
    for r in range(2, ws.max_row + 1):
        reg = str(ws.cell(row=r, column=1).value or '').strip().upper()
        subj = str(ws.cell(row=r, column=3).value or '').strip().upper()
        test = str(ws.cell(row=r, column=4).value or '').strip()
        if reg and subj and test:
            existing_key_map[(reg, subj, test)] = r

    # Process and upsert each record
    for rec in records:
        reg_no = str(rec.get('regno') or '').strip()
        name = str(rec.get('student_name') or '').strip()
        subj = str(rec.get('subject_name') or '').strip()
        test_num = str(rec.get('test_number') or '1').strip()
        max_marks = float(rec.get('max_marks') or 100.0)

        # Marks obtained: display 'Absent' if absent or None, otherwise numeric
        if rec.get('is_absent') or rec.get('marks_obtained') is None:
            marks_obt = "Absent"
        else:
            marks_obt = float(rec.get('marks_obtained'))

        lookup_key = (reg_no.upper(), subj.upper(), test_num)

        if lookup_key in existing_key_map:
            # Update existing row
            target_row = existing_key_map[lookup_key]
            ws.cell(row=target_row, column=1, value=reg_no).data_type = 's'
            ws.cell(row=target_row, column=2, value=name)
            ws.cell(row=target_row, column=3, value=subj)
            ws.cell(row=target_row, column=4, value=int(test_num) if test_num.isdigit() else test_num)
            ws.cell(row=target_row, column=5, value=marks_obt)
            ws.cell(row=target_row, column=6, value=max_marks)
        else:
            # Append new row
            new_row = ws.max_row + 1
            ws.cell(row=new_row, column=1, value=reg_no).data_type = 's'
            ws.cell(row=new_row, column=2, value=name)
            ws.cell(row=new_row, column=3, value=subj)
            ws.cell(row=new_row, column=4, value=int(test_num) if test_num.isdigit() else test_num)
            ws.cell(row=new_row, column=5, value=marks_obt)
            ws.cell(row=new_row, column=6, value=max_marks)
            existing_key_map[lookup_key] = new_row

    # Apply professional styling
    _apply_table_styling(ws)

    # Save to primary and mirror paths
    wb.save(primary_path)
    if mirror_path:
        try:
            wb.save(mirror_path)
        except Exception as e:
            print(f"Notice: Mirror marks_data.xlsx save failed ({e})")


def regenerate_full_marks_data_file():
    """
    Pulls all internal marks records from the live database and populates
    marks_data.xlsx in sheet 'Marks' from scratch.
    """
    query = """
    SELECT s.regno, s.name as student_name, sub.subject_name,
           m.test_number, m.marks_obtained, m.max_marks
    FROM internal_marks m
    JOIN students s ON m.student_id = s.studentid
    JOIN subjects sub ON m.subject_id = sub.subjectid
    ORDER BY sub.subject_name ASC, m.test_number ASC, s.regno ASC
    """
    rows = fetch_all(query) or []
    
    primary_path, mirror_path = get_marks_data_filepaths()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Marks"
    
    headers = ["Register No", "Student Name", "Subject", "Test Number", "Marks Obtained", "Max Marks"]
    ws.append(headers)

    for r in rows:
        reg_no = str(r.get('regno') or '')
        name = str(r.get('student_name') or '')
        subj = str(r.get('subject_name') or '')
        test_num = int(r.get('test_number') or 1)
        m_obt = r.get('marks_obtained')
        max_m = float(r.get('max_marks') or 100.0)
        
        marks_val = "Absent" if m_obt is None else float(m_obt)

        row_num = ws.max_row + 1
        ws.cell(row=row_num, column=1, value=reg_no).data_type = 's'
        ws.cell(row=row_num, column=2, value=name)
        ws.cell(row=row_num, column=3, value=subj)
        ws.cell(row=row_num, column=4, value=test_num)
        ws.cell(row=row_num, column=5, value=marks_val)
        ws.cell(row=row_num, column=6, value=max_m)

    _apply_table_styling(ws)
    wb.save(primary_path)
    if mirror_path:
        try:
            wb.save(mirror_path)
        except Exception:
            pass

    return primary_path
